"""Atomically generate and publish OKF round-trip fixtures.

All source binaries, frozen Docling node output, expected direct-chain spans,
and fixture READMEs are generated beneath one same-volume staging directory.
Every staged fixture is validated before any published artifact changes. Only
then are complete fixture directories swapped into place with rollback.

Run only with the isolated Docling 2.109 environment, from ``<repo-root>``::

    C:/Users/daixu/AppData/Local/Temp/docling-hierarchy-verify-2109/venv/Scripts/python.exe \\
      tests/fixtures/okf_roundtrip/generate_fixtures.py
"""

from __future__ import annotations

# Standalone script: repository and sibling module paths are added below before
# local imports, so module-import ordering cannot remain at the lexical top.
# ruff: noqa: E402

import json
import os
import shutil
import sys
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import NAMESPACE_URL, UUID, uuid4, uuid5

_REPO_ROOT = Path(__file__).resolve().parents[3]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from llamaindex_runtime.ingestion.docling_ingestor import DoclingIngestor
from llamaindex_runtime.ingestion.normalization import NormalizationContract
from validate_fixtures import validate_fixture

FIXTURE_DOC_ID = "00000000-0000-0000-0000-000000000001"
FIXTURE_VERSION_ID = "00000000-0000-0000-0000-000000000003"
FIXTURE_DIR = Path(__file__).parent
FIXTURE_NAMES = ("sectioned-pdf", "complex-layout-pdf", "docx")
REQUIRED_DOCLING_VERSION = "2.109.0"


class ValidationError(Exception):
    """A generated fixture did not meet its required contract."""


class RecoveryError(Exception):
    """Rollback failed during publication; recovery dirs are preserved.

    Raised when the rollback phase of ``publish_staged_artifacts`` itself
    encounters an error. The ``backup_root`` and/or ``stage_root`` directories
    are deliberately NOT cleaned up so the operator can recover original
    artifacts manually. The error message names the preserved paths.
    """

    def __init__(
        self,
        message: str,
        backup_root: Path | None = None,
        stage_root: Path | None = None,
    ) -> None:
        self.backup_root = backup_root
        self.stage_root = stage_root
        paths = []
        if backup_root:
            paths.append(f"backup={backup_root}")
        if stage_root:
            paths.append(f"stage={stage_root}")
        full_msg = f"{message} Preserved recovery dirs: {', '.join(paths)}"
        super().__init__(full_msg)


def assert_docling_version() -> None:
    """Assert the runtime docling version equals 2.109.0.

    The generator and production reader depend on HeadingHierarchyOptions
    which is only available in docling 2.109.x. This guard ensures the
    FIXTURE_PREPARATION_STATUS claim is truthful at generation time.
    """
    import importlib.metadata

    actual = importlib.metadata.version("docling")
    if actual != REQUIRED_DOCLING_VERSION:
        raise ValidationError(
            f"docling version must be {REQUIRED_DOCLING_VERSION}, got {actual}"
        )


class HeadingOutlineDocTemplate(BaseDocTemplate):
    """Attach PDF outline destinations at the actual heading flow location."""

    _OUTLINE_LEVELS = {"FixtureH1": 0, "FixtureH2": 1, "FixtureH3": 2}

    def __init__(self, buffer: BytesIO) -> None:
        super().__init__(buffer, pagesize=letter, invariant=1)
        frame = Frame(
            self.leftMargin, self.bottomMargin, self.width, self.height, id="body"
        )
        self.addPageTemplates([PageTemplate(id="main", frames=[frame])])
        self._bookmark_counter = 0

    def afterFlowable(self, flowable: object) -> None:
        """Bookmark each heading when ReportLab has actually positioned it."""
        style_name = getattr(getattr(flowable, "style", None), "name", "")
        level = self._OUTLINE_LEVELS.get(style_name)
        if level is None:
            return
        text = getattr(flowable, "getPlainText", lambda: "")()
        if not text:
            return
        self._bookmark_counter += 1
        destination = f"heading-{self._bookmark_counter}"
        self.canv.bookmarkPage(destination)
        self.canv.addOutlineEntry(text, destination, level=level)


def _styles() -> tuple[ParagraphStyle, ParagraphStyle, ParagraphStyle, ParagraphStyle]:
    styles = getSampleStyleSheet()
    return (
        ParagraphStyle(
            "FixtureH1", parent=styles["Heading1"], fontSize=18, spaceAfter=12
        ),
        ParagraphStyle(
            "FixtureH2", parent=styles["Heading2"], fontSize=14, spaceAfter=10
        ),
        ParagraphStyle(
            "FixtureH3", parent=styles["Heading3"], fontSize=12, spaceAfter=8
        ),
        ParagraphStyle(
            "FixtureBody", parent=styles["Normal"], fontSize=10, spaceAfter=6
        ),
    )


_PDF_FONT_CACHE_PRIMED = False


def _build_pdf(story: list[object]) -> bytes:
    """Build a deterministic PDF.

    ReportLab's first build in a process embeds fonts that subsequent builds
    omit from the output (the style registry caches them).  A throwaway prime
    build with a minimal dummy story ensures every real build produces
    identical bytes.  The dummy story is used instead of the real story because
    ReportLab mutates flowable state during build.
    """
    global _PDF_FONT_CACHE_PRIMED
    if not _PDF_FONT_CACHE_PRIMED:
        dummy = [Paragraph("prime", _styles()[3])]
        HeadingOutlineDocTemplate(BytesIO()).build(dummy)
        _PDF_FONT_CACHE_PRIMED = True
    buffer = BytesIO()
    HeadingOutlineDocTemplate(buffer).build(story)
    return buffer.getvalue()


def create_sectioned_pdf() -> bytes:
    """Create a real PDF whose hierarchy bookmarks flow with their headings."""
    h1, h2, h3, body = _styles()
    return _build_pdf(
        [
            Paragraph("Introduction to Machine Learning", h1),
            Paragraph("Overview of ML fundamentals.", body),
            Paragraph("Chapter 1: Foundations", h2),
            Paragraph(
                "Machine learning is a subset of artificial intelligence that enables systems to learn from data.",
                body,
            ),
            Paragraph("1.1 Supervised Learning", h3),
            Paragraph(
                "Supervised learning uses labeled training data to learn mapping functions.",
                body,
            ),
            Paragraph("1.2 Unsupervised Learning", h3),
            Paragraph(
                "Unsupervised learning discovers hidden patterns in unlabeled data.",
                body,
            ),
            Paragraph("Chapter 2: Applications", h2),
            Paragraph("Machine learning applications span many industries.", body),
            Paragraph("2.1 Natural Language Processing", h3),
            Paragraph(
                "NLP enables computers to understand and process human language.", body
            ),
        ]
    )


def create_complex_layout_pdf() -> bytes:
    """Create a table PDF whose heading bookmarks flow with their headings."""
    h1, h2, h3, body = _styles()
    table = Table(
        [
            ["Category", "Q3 2024", "Q4 2024", "Change"],
            ["Revenue", "12,500", "14,200", "+13.6%"],
            ["Expenses", "8,200", "8,900", "+8.5%"],
            ["Profit", "4,300", "5,300", "+23.3%"],
        ],
        colWidths=[1.5 * inch, inch, inch, inch],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]
        )
    )
    return _build_pdf(
        [
            Paragraph("Quarterly Financial Report", h1),
            Paragraph(
                "This report summarizes Q4 2024 financial performance metrics.", body
            ),
            table,
            Spacer(1, 12),
            Paragraph("Analysis", h2),
            Paragraph(
                "The strong Q4 performance reflects improved market conditions.", body
            ),
            Paragraph("Regional Breakdown", h3),
            Paragraph("North America contributed 45% of total revenue.", body),
        ]
    )


_CANONICAL_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def _canonicalize_docx_zip(data: bytes) -> bytes:
    """Re-pack a DOCX ZIP with deterministic entry timestamps.

    python-docx delegates to ``zipfile.ZipFile.writestr`` which stamps every
    entry with the current wall-clock ``date_time``.  This function re-writes
    the entire ZIP with a fixed canonical timestamp ``(1980, 1, 1, 0, 0, 0)``
    (the MS-DOS epoch minimum) so that two generator runs produce
    byte-identical DOCX files.

    The following fields are explicitly copied from each original ``ZipInfo``
    to the replacement: ``compress_type``, ``external_attr``,
    ``internal_attr``, ``create_system``, ``extract_version``,
    ``create_version``, ``flag_bits``.  Entry ordering, filenames, and
    decompressed content bytes are preserved exactly.  Only the timestamp
    (``date_time``) is canonicalized.  The result remains a valid OPC/DOCX
    that python-docx and Docling can read.

    Both source and destination ZIPs are opened with context managers so
    resources are released even on exception.
    """
    output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(data)) as source_zip,
        zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as dest_zip,
    ):
        for info in source_zip.infolist():
            new_info = zipfile.ZipInfo(
                filename=info.filename,
                date_time=_CANONICAL_ZIP_TIMESTAMP,
            )
            new_info.compress_type = info.compress_type
            new_info.external_attr = info.external_attr
            new_info.internal_attr = info.internal_attr
            new_info.create_system = info.create_system
            new_info.extract_version = info.extract_version
            new_info.create_version = info.create_version
            new_info.flag_bits = info.flag_bits
            dest_zip.writestr(new_info, source_zip.read(info.filename))
    return output.getvalue()


def create_unicode_docx() -> bytes:
    """Create a DOCX with CJK text and emoji for the non-PDF path."""
    from docx import Document

    document = Document()
    document.add_heading("全球产品公告 / Global Product Announcement", level=1)
    document.add_heading("国际市场 / International Markets", level=2)
    document.add_paragraph(
        "Our product is now available worldwide: Americas, Europe, Asia-Pacific."
    )
    document.add_heading("中国市场 / Chinese Market", level=3)
    document.add_paragraph(
        "我们很高兴地宣布进入中国市场。 We are excited to announce expansion."
    )
    document.add_heading("特殊字符测试 / Special Characters Test", level=2)
    document.add_paragraph("Emoji: 🌐🌍🌎🕊 Symbols: ☕❤✓ Math: Σ ≠ ≤ ≥ π.")
    buffer = BytesIO()
    document.save(buffer)
    return _canonicalize_docx_zip(buffer.getvalue())


def _freeze_node(node: object) -> dict[str, Any]:
    return {
        "text": DoclingIngestor._extract_text(node),
        "metadata": dict(getattr(node, "metadata", {}) or {}),
    }


def _direct_spans(
    nodes: list[object], doc_id: str, version_id: str
) -> list[dict[str, Any]]:
    """Compute frozen coordinates with the exact production normalization chain."""
    contract = NormalizationContract()
    spans: list[dict[str, Any]] = []
    for ordinal, node in enumerate(nodes):
        raw_text = DoclingIngestor._extract_text(node)
        if not raw_text.strip():
            continue
        normalized = contract.normalize(
            raw_text=raw_text,
            metadata=DoclingIngestor._flatten_docling_metadata(
                dict(getattr(node, "metadata", {}) or {})
            ),
            ordinal=ordinal,
        )
        span_id = uuid5(
            NAMESPACE_URL,
            f"{doc_id}|{version_id}|{normalized.page_no}|{'/'.join(normalized.headings)}|"
            f"{normalized.offset}|{normalized.text}",
        )
        spans.append(
            {
                "span_id": str(span_id),
                "page_no": normalized.page_no,
                "heading_path": list(normalized.headings),
                "offset": normalized.offset,
                "text": normalized.text,
            }
        )
    return spans


def _fixture_readme(fixture_name: str, source_ext: str) -> str:
    description = {
        "sectioned-pdf": "PDF with multi-level heading structure.",
        "complex-layout-pdf": "PDF with a table, complex layout, and heading hierarchy.",
        "docx": "DOCX with Unicode, Chinese, and emoji content.",
    }[fixture_name]
    return f"""# {fixture_name} Fixture

## Description

{description}

## Contents

| File | Description |
|------|-------------|
| source.{source_ext} | Locally authored source document |
| docling_output.json | Frozen `DoclingNodeParser` node sequence |
| expected_span_ids.json | Expected coordinates and span IDs from the production direct chain |

## Fixed UUIDs

- doc_id: `{FIXTURE_DOC_ID}`
- version_id: `{FIXTURE_VERSION_ID}`

## Regeneration

Run from `<repo-root>` with the required isolated environment:

```bash
C:/Users/daixu/AppData/Local/Temp/docling-hierarchy-verify-2109/venv/Scripts/python.exe \\
  tests/fixtures/okf_roundtrip/generate_fixtures.py
```

The generator stages all three fixtures, validates all staged outputs, then
atomically swaps complete fixture directories into place. A failure leaves the
published fixture set unchanged when rollback succeeds. If rollback fails, the
generator preserves named stage and backup directories for manual recovery.

**Docling version**: `2.109.0`

## Migration characterization

The hierarchy reader intentionally changes `heading_path` and therefore
`span_id` for aligned nested PDF spans. The production migration test proves
that matching spans retain `page_no`, `offset`, and `text`; root headings can
remain unchanged while nested headings gain hierarchy. A custom reader can be
injected into `DoclingIngestor` for rollback.

## License

Locally authored. Apache-2.0 compatible.
"""


def _main_readme() -> str:
    return """# OKF Round-Trip Fixtures

| Fixture | Format | Key Test |
|---------|--------|----------|
| sectioned-pdf | PDF | Multi-level heading_path |
| complex-layout-pdf | PDF | Tables and nested structure |
| docx | DOCX | Unicode/Chinese/emoji |

## Validation

From `<repo-root>`:

```bash
python tests/fixtures/okf_roundtrip/validate_fixtures.py
```

## Regeneration

```bash
C:/Users/daixu/AppData/Local/Temp/docling-hierarchy-verify-2109/venv/Scripts/python.exe \\
  tests/fixtures/okf_roundtrip/generate_fixtures.py
```
"""


def generate_fixture(
    stage_root: Path,
    fixture_name: str,
    source_bytes: bytes,
    source_ext: str,
) -> None:
    """Create one complete fixture beneath ``stage_root``; never publish it."""
    fixture_dir = stage_root / fixture_name
    fixture_dir.mkdir()
    source_path = fixture_dir / f"source.{source_ext}"
    source_path.write_bytes(source_bytes)

    ingestor = DoclingIngestor()
    documents = ingestor._reader.load_data(file_path=str(source_path))
    nodes = list(ingestor._node_parser.get_nodes_from_documents(documents))
    if not nodes:
        raise ValidationError(f"[{fixture_name}] Docling extracted no nodes")

    frozen_nodes = [_freeze_node(node) for node in nodes]
    spans = _direct_spans(nodes, FIXTURE_DOC_ID, FIXTURE_VERSION_ID)
    if not spans:
        raise ValidationError(f"[{fixture_name}] direct chain produced no spans")

    # Cross-check the frozen direct computation against the public production
    # entry point before anything can be published.
    direct_result = ingestor.ingest(
        source_path, doc_id=UUID(FIXTURE_DOC_ID), version_id=UUID(FIXTURE_VERSION_ID)
    )
    direct_records = [
        {
            "span_id": str(span.span_id),
            "page_no": span.page_no,
            "heading_path": list(span.headings),
            "offset": span.offset,
            "text": span.text,
        }
        for span in direct_result.spans
    ]
    if direct_records != spans:
        raise ValidationError(
            f"[{fixture_name}] production ingest differs from frozen direct chain"
        )

    (fixture_dir / "docling_output.json").write_text(
        json.dumps(frozen_nodes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (fixture_dir / "expected_span_ids.json").write_text(
        json.dumps(
            {
                "doc_id": FIXTURE_DOC_ID,
                "version_id": FIXTURE_VERSION_ID,
                "spans": spans,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    (fixture_dir / "README.md").write_text(
        _fixture_readme(fixture_name, source_ext), encoding="utf-8"
    )


def _remove_path(path: Path) -> None:
    if path.is_dir():
        shutil.rmtree(path)
    elif path.exists():
        path.unlink()


def publish_staged_artifacts(
    stage_root: Path,
    fixture_root: Path,
    *,
    fixture_names: tuple[str, ...] = FIXTURE_NAMES,
) -> None:
    """Atomically swap staged fixture directories, rolling back on any failure.

    ``stage_root`` must be created under ``fixture_root.parent`` so every
    ``os.replace`` is same-volume. Complete fixture directories are the atomic
    unit; the parent generator/validator/test scripts are deliberately never
    replaced or deleted. The main README is published only after all fixture
    directory swaps succeed.

    Rollback safety:
        If a publish step fails, every already-swapped artifact is restored
        from backup. If rollback itself fails, a ``RecoveryError`` is raised
        and the backup/stage directories are preserved for manual recovery.
        Cleanup of transient directories only happens when publish succeeds or
        rollback fully succeeds.
    """
    fixture_root = fixture_root.resolve()
    stage_root = stage_root.resolve()
    if stage_root.parent != fixture_root.parent:
        raise ValueError(
            "stage_root must share fixture_root's parent for same-volume swaps"
        )
    for name in fixture_names:
        if not (stage_root / name).is_dir():
            raise ValidationError(f"staged fixture directory missing: {name}")
    staged_readme = stage_root / "README.md"
    if not staged_readme.is_file():
        raise ValidationError("staged main README missing")

    backups: list[tuple[Path, Path]] = []
    published: list[tuple[Path, Path]] = []
    backup_root = fixture_root.parent / f".{fixture_root.name}.backup-{uuid4().hex}"
    backup_root.mkdir()
    rollback_failed = False

    try:
        for name in fixture_names:
            target = fixture_root / name
            staged = stage_root / name
            backup = backup_root / name
            if target.exists():
                os.replace(target, backup)
                backups.append((target, backup))
            os.replace(staged, target)
            published.append((target, staged))

        # Publish the main README last, only after every fixture validates and
        # every fixture directory has swapped successfully.
        target_readme = fixture_root / "README.md"
        backup_readme = backup_root / "README.md"
        if target_readme.exists():
            os.replace(target_readme, backup_readme)
            backups.append((target_readme, backup_readme))
        os.replace(staged_readme, target_readme)
        published.append((target_readme, staged_readme))
    except Exception:
        # Rollback: move newly published artifacts back to staged, then
        # restore originals from backup.
        try:
            for target, staged in reversed(published):
                if target.exists():
                    os.replace(target, staged)
            for target, backup in reversed(backups):
                if backup.exists():
                    os.replace(backup, target)
        except Exception:
            rollback_failed = True
            raise RecoveryError(
                "Rollback failed during publish recovery. "
                "Original artifacts may be in backup or staged dirs.",
                backup_root=backup_root,
                stage_root=stage_root,
            ) from None
        raise
    finally:
        # Only clean up transients when everything succeeded or rollback
        # completed without error. If rollback failed, preserve dirs for
        # manual recovery.
        if not rollback_failed:
            _remove_path(stage_root)
            _remove_path(backup_root)


def main() -> int:
    """Stage all fixtures, validate all three, then publish them together."""
    assert_docling_version()
    stage_root = Path(
        tempfile.mkdtemp(prefix=".okf-roundtrip-stage-", dir=FIXTURE_DIR.parent)
    )
    recovery_required = False
    try:
        generate_fixture(stage_root, "sectioned-pdf", create_sectioned_pdf(), "pdf")
        generate_fixture(
            stage_root, "complex-layout-pdf", create_complex_layout_pdf(), "pdf"
        )
        generate_fixture(stage_root, "docx", create_unicode_docx(), "docx")
        (stage_root / "README.md").write_text(_main_readme(), encoding="utf-8")

        failures = [
            result
            for name in FIXTURE_NAMES
            if not (result := validate_fixture(name, fixture_root=stage_root))["valid"]
        ]
        if failures:
            details = "; ".join(
                f"{result['name']}: {', '.join(result['errors'])}"
                for result in failures
            )
            raise ValidationError(f"staged fixture validation failed: {details}")

        publish_staged_artifacts(stage_root, FIXTURE_DIR)
        print("ALL FIXTURES GENERATED, VALIDATED, AND ATOMICALLY PUBLISHED")
        return 0
    except (ValidationError, OSError) as exc:
        print(f"FIXTURE GENERATION FAILED: {exc}", file=sys.stderr)
        return 1
    except RecoveryError as exc:
        recovery_required = True
        print(f"FIXTURE GENERATION RECOVERY REQUIRED: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"FIXTURE GENERATION FAILED unexpectedly: {exc}", file=sys.stderr)
        return 1
    finally:
        # Conversion/validation/publish failures must not leave stale staged
        # artifacts behind. A failed rollback deliberately preserves the
        # recovery roots named by ``RecoveryError``.
        if not recovery_required:
            _remove_path(stage_root)


if __name__ == "__main__":
    raise SystemExit(main())

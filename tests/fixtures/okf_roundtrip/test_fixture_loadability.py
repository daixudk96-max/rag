"""Test that fixtures can be loaded and processed.

Validates:
1. JSON parsing
2. UUID validity
3. heading_path depth requirements
4. Unicode content requirements
5. Recomputation matches expected
6. Source builder byte-determinism
"""

# Standalone test module: it bootstraps sibling/repository paths below before
# imports, so lexical import ordering cannot be preserved.
# ruff: noqa: E402
from __future__ import annotations

import hashlib
import json
import os
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import UUID, NAMESPACE_URL, uuid5

import pytest

_FIXTURE_MODULE_DIR = Path(__file__).parent
if str(_FIXTURE_MODULE_DIR) not in sys.path:
    sys.path.insert(0, str(_FIXTURE_MODULE_DIR))
import generate_fixtures

sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))

from llamaindex_runtime.ingestion.docling_ingestor import DoclingIngestor
from llamaindex_runtime.ingestion.normalization import NormalizationContract

FIXTURE_DIR = (
    Path(__file__).parent.parent.parent.parent / "tests" / "fixtures" / "okf_roundtrip"
)


def load_fixture(fixture_name: str) -> tuple[list[dict], dict]:
    """Load fixture files."""
    fixture_dir = FIXTURE_DIR / fixture_name

    with open(fixture_dir / "docling_output.json", encoding="utf-8") as f:
        nodes = json.load(f)

    with open(fixture_dir / "expected_span_ids.json", encoding="utf-8") as f:
        expected = json.load(f)

    return nodes, expected


class TestFixtureLoadability:
    """Verify all fixtures load correctly."""

    def test_sectioned_loads(self):
        """Sectioned fixture loads."""
        nodes, expected = load_fixture("sectioned-pdf")
        assert len(nodes) > 0
        assert len(expected["spans"]) > 0

    def test_complex_loads(self):
        """Complex fixture loads."""
        nodes, expected = load_fixture("complex-layout-pdf")
        assert len(nodes) > 0
        assert len(expected["spans"]) > 0

    def test_docx_loads(self):
        """DOCX fixture loads."""
        nodes, expected = load_fixture("docx")
        assert len(nodes) > 0
        assert len(expected["spans"]) > 0


class TestFixtureUuids:
    """Verify all UUIDs are valid."""

    @pytest.mark.parametrize(
        "fixture_name", ["sectioned-pdf", "complex-layout-pdf", "docx"]
    )
    def test_uuids_valid(self, fixture_name: str):
        """All UUIDs are valid."""
        _, expected = load_fixture(fixture_name)

        UUID(expected["doc_id"])
        UUID(expected["version_id"])

        for i, span in enumerate(expected["spans"]):
            UUID(span["span_id"])


class TestFixtureSemantics:
    """Verify semantic requirements."""

    def test_sectioned_requires_multilevel_headings(self):
        """The sectioned PDF hard gate requires a depth-two heading path."""
        _, expected = load_fixture("sectioned-pdf")

        max_depth = max(len(span["heading_path"]) for span in expected["spans"])
        assert max_depth >= 2, f"Expected heading depth >= 2, got {max_depth}"

    def test_complex_has_table_provenance_and_charspan(self):
        """Complex PDF preserves a table node with PDF provenance and charspan."""
        nodes, _ = load_fixture("complex-layout-pdf")

        assert any(
            node["metadata"].get("doc_items")
            and node["metadata"]["doc_items"][0].get("label") == "table"
            and node["metadata"]["doc_items"][0].get("prov")
            and "page_no" in node["metadata"]["doc_items"][0]["prov"][0]
            and "charspan" in node["metadata"]["doc_items"][0]["prov"][0]
            for node in nodes
        ), "Expected a table doc_item with provenance page_no and charspan"

    def test_docx_unicode_content(self):
        """DOCX fixture has actual Chinese characters and emoji."""
        _, expected = load_fixture("docx")

        all_text = " ".join(span["text"] for span in expected["spans"])
        all_headings = " ".join(
            h for span in expected["spans"] for h in span["heading_path"]
        )

        # Check for Chinese characters (hard gate)
        has_chinese = any("一" <= c <= "鿿" for c in all_text + all_headings)
        assert has_chinese, "Expected Chinese characters in docx fixture"

        # Check for emoji (hard gate)
        has_emoji = any(ord(c) > 0x1F000 for c in all_text + all_headings)
        assert has_emoji, "Expected emoji in docx fixture"

    def test_complex_has_table_content(self):
        """Complex fixture has table-related content."""
        _, expected = load_fixture("complex-layout-pdf")

        all_text = " ".join(span["text"] for span in expected["spans"]).lower()

        # Table content should have financial keywords
        table_keywords = ["revenue", "expenses", "profit", "q3", "q4"]
        has_table = any(kw in all_text for kw in table_keywords)
        assert has_table, "Expected table content in complex-layout fixture"


class TestRecomputation:
    """Verify recomputation from frozen nodes matches expected.

    Compares full coordinates: span_id, page_no, heading_path, offset, text and order.
    Does NOT invoke Docling conversion - frozen JSON tests run fully offline.
    """

    def recompute_span_id(
        self,
        doc_id: str,
        version_id: str,
        page_no: int | None,
        heading_path: list[str],
        offset: int,
        text: str,
    ) -> str:
        """Recompute span_id using exact formula."""
        return str(
            uuid5(
                NAMESPACE_URL,
                f"{doc_id}|{version_id}|{page_no}|{'/'.join(heading_path)}|{offset}|{text}",
            )
        )

    @pytest.mark.parametrize(
        "fixture_name", ["sectioned-pdf", "complex-layout-pdf", "docx"]
    )
    def test_recomputation_matches(self, fixture_name: str):
        """Recomputed spans match expected using EXACT direct chain.

        Compares span_id, page_no, heading_path, offset, text in document order.
        """
        nodes, expected = load_fixture(fixture_name)

        doc_id = expected["doc_id"]
        version_id = expected["version_id"]

        # EXACT direct chain from generate_fixtures.py
        contract = NormalizationContract()
        recomputed = []

        for ordinal, node in enumerate(nodes):
            text = node.get("text", "")
            if not text.strip():
                continue

            metadata = node.get("metadata", {})
            # EXACT DoclingIngestor._flatten_docling_metadata
            flat_metadata = DoclingIngestor._flatten_docling_metadata(metadata)

            # EXACT NormalizationContract.normalize
            normalized = contract.normalize(
                raw_text=text,
                metadata=flat_metadata,
                ordinal=ordinal,
            )

            span_id = str(
                uuid5(
                    NAMESPACE_URL,
                    f"{doc_id}|{version_id}|{normalized.page_no}|"
                    f"{'/'.join(normalized.headings)}|{normalized.offset}|{normalized.text}",
                )
            )

            recomputed.append(
                {
                    "span_id": span_id,
                    "page_no": normalized.page_no,
                    "heading_path": list(normalized.headings),
                    "offset": normalized.offset,
                    "text": normalized.text,
                }
            )

        expected_spans = expected["spans"]

        # Compare count
        assert len(recomputed) == len(expected_spans), (
            f"[{fixture_name}] span count mismatch: "
            f"recomputed={len(recomputed)}, expected={len(expected_spans)}"
        )

        # Compare full coordinates in order
        for i, (rec, exp) in enumerate(zip(recomputed, expected_spans)):
            for field in ("span_id", "page_no", "heading_path", "offset", "text"):
                assert rec[field] == exp[field], (
                    f"[{fixture_name}] span[{i}] {field} mismatch: "
                    f"recomputed={rec[field]!r}, expected={exp[field]!r}"
                )


class TestAtomicFixturePublication:
    """Exercise publication rollback without invoking Docling conversion."""

    def test_publish_failure_restores_existing_fixture_tree(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A failed swap with successful rollback must restore byte-for-byte.

        Injecting a same-volume rename failure on the second fixture proves
        that an earlier fixture swap is rolled back and that no stale
        stage/backup directories remain.
        """
        fixture_root = tmp_path / "okf_roundtrip"
        fixture_root.mkdir()
        names = ("first", "second")
        for name in names:
            target = fixture_root / name
            target.mkdir()
            (target / "payload.txt").write_text(f"old-{name}", encoding="utf-8")
        (fixture_root / "README.md").write_text("old README", encoding="utf-8")

        staged_root = fixture_root.parent / ".okf-roundtrip-stage-test"
        staged_root.mkdir()
        for name in names:
            staged = staged_root / name
            staged.mkdir()
            (staged / "payload.txt").write_text(f"new-{name}", encoding="utf-8")
        (staged_root / "README.md").write_text("new README", encoding="utf-8")

        real_replace = os.replace

        def fail_on_second(source: str | Path, destination: str | Path) -> None:
            if Path(source).name == "second":
                raise OSError("injected publication failure")
            real_replace(source, destination)

        monkeypatch.setattr(generate_fixtures.os, "replace", fail_on_second)

        with pytest.raises(OSError, match="injected publication failure"):
            generate_fixtures.publish_staged_artifacts(
                staged_root, fixture_root, fixture_names=names
            )

        assert (fixture_root / "first" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "old-first"
        assert (fixture_root / "second" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "old-second"
        assert (fixture_root / "README.md").read_text(encoding="utf-8") == "old README"
        assert not staged_root.exists()
        assert not list(fixture_root.glob(".*.backup-*"))

    def test_rollback_failure_preserves_recovery_dirs(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """If rollback itself fails, recovery dirs must be preserved.

        CRITICAL scenario: publish fails after original target/second has been
        moved to backup. Rollback attempts to restore backup -> target but that
        restore also fails. The publisher must raise a dedicated recovery error
        containing the preserved recovery paths. The backup and staged dirs
        must NOT be cleaned up because they hold the only recoverable copies.
        """
        fixture_root = tmp_path / "okf_roundtrip"
        fixture_root.mkdir()
        names = ("first", "second")
        for name in names:
            target = fixture_root / name
            target.mkdir()
            (target / "payload.txt").write_text(f"old-{name}", encoding="utf-8")
        (fixture_root / "README.md").write_text("old README", encoding="utf-8")

        staged_root = fixture_root.parent / ".okf-roundtrip-stage-test"
        staged_root.mkdir()
        for name in names:
            staged = staged_root / name
            staged.mkdir()
            (staged / "payload.txt").write_text(f"new-{name}", encoding="utf-8")
        (staged_root / "README.md").write_text("new README", encoding="utf-8")

        real_replace = os.replace

        def controlled_replace(source: str | Path, destination: str | Path) -> None:
            src = Path(source)
            dst = Path(destination)
            is_second_publish = (
                src == staged_root / "second" and dst == fixture_root / "second"
            )
            is_first_backup_restore = (
                src.parent.name.startswith(f".{fixture_root.name}.backup-")
                and src.name == "first"
                and dst == fixture_root / "first"
            )
            if is_second_publish:
                raise OSError("injected publish failure on second")
            if is_first_backup_restore:
                raise OSError("injected rollback failure")
            real_replace(source, destination)

        monkeypatch.setattr(generate_fixtures.os, "replace", controlled_replace)

        with pytest.raises(generate_fixtures.RecoveryError) as exc_info:
            generate_fixtures.publish_staged_artifacts(
                staged_root, fixture_root, fixture_names=names
            )

        error_msg = str(exc_info.value)
        # The recovery error must name the preserved recovery directories
        assert (
            "backup" in error_msg.lower() or "recovery" in error_msg.lower()
        ), f"RecoveryError must mention backup/recovery paths, got: {error_msg}"

        # The backup directory must still exist with recoverable original bytes
        backup_dirs = list(fixture_root.parent.glob(f".{fixture_root.name}.backup-*"))
        assert (
            len(backup_dirs) == 1
        ), f"Expected exactly 1 preserved backup dir, found {len(backup_dirs)}"
        backup_dir = backup_dirs[0]
        assert (backup_dir / "first" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "old-first", "Original bytes must be recoverable in backup dir"

        # The staged directory may or may not exist depending on where rollback
        # failed, but at least one recovery directory must survive cleanup.
        # The critical invariant: backup must NOT be deleted.

    def test_successful_publish_cleans_all_transients(self, tmp_path: Path) -> None:
        """On fully successful publish, no backup or stage dirs remain."""
        fixture_root = tmp_path / "okf_roundtrip"
        fixture_root.mkdir()
        names = ("first", "second")
        for name in names:
            target = fixture_root / name
            target.mkdir()
            (target / "payload.txt").write_text(f"old-{name}", encoding="utf-8")
        (fixture_root / "README.md").write_text("old README", encoding="utf-8")

        staged_root = fixture_root.parent / ".okf-roundtrip-stage-test"
        staged_root.mkdir()
        for name in names:
            staged = staged_root / name
            staged.mkdir()
            (staged / "payload.txt").write_text(f"new-{name}", encoding="utf-8")
        (staged_root / "README.md").write_text("new README", encoding="utf-8")

        generate_fixtures.publish_staged_artifacts(
            staged_root, fixture_root, fixture_names=names
        )

        assert (fixture_root / "first" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "new-first"
        assert (fixture_root / "second" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "new-second"
        assert (fixture_root / "README.md").read_text(encoding="utf-8") == "new README"
        assert not staged_root.exists()
        assert not list(fixture_root.parent.glob(f".{fixture_root.name}.backup-*"))

    def test_main_preserves_both_recovery_roots_after_incomplete_rollback(
        self,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        """The CLI retains the recovery paths named by RecoveryError."""
        fixture_root = tmp_path / "okf_roundtrip"
        fixture_root.mkdir()
        for name in generate_fixtures.FIXTURE_NAMES:
            target = fixture_root / name
            target.mkdir()
            (target / "payload.txt").write_text(f"old-{name}", encoding="utf-8")
        (fixture_root / "README.md").write_text("old README", encoding="utf-8")

        monkeypatch.setattr(generate_fixtures, "FIXTURE_DIR", fixture_root)
        monkeypatch.setattr(generate_fixtures, "assert_docling_version", lambda: None)
        monkeypatch.setattr(generate_fixtures, "create_sectioned_pdf", lambda: b"pdf")
        monkeypatch.setattr(
            generate_fixtures, "create_complex_layout_pdf", lambda: b"pdf"
        )
        monkeypatch.setattr(generate_fixtures, "create_unicode_docx", lambda: b"docx")
        monkeypatch.setattr(
            generate_fixtures,
            "validate_fixture",
            lambda *args, **kwargs: {"valid": True},
        )

        def generate_staged_fixture(
            stage_root: Path,
            fixture_name: str,
            source_bytes: bytes,
            source_ext: str,
        ) -> None:
            fixture_dir = stage_root / fixture_name
            fixture_dir.mkdir()
            (fixture_dir / "payload.txt").write_text(
                f"new-{fixture_name}", encoding="utf-8"
            )

        monkeypatch.setattr(
            generate_fixtures, "generate_fixture", generate_staged_fixture
        )
        real_replace = os.replace

        def fail_publish_then_rollback(
            source: str | Path, destination: str | Path
        ) -> None:
            src = Path(source)
            dst = Path(destination)
            is_second_publish = (
                src.parent.name.startswith(".okf-roundtrip-stage-")
                and src.name == "complex-layout-pdf"
                and dst == fixture_root / "complex-layout-pdf"
            )
            is_first_backup_restore = (
                src.parent.name.startswith(f".{fixture_root.name}.backup-")
                and src.name == "sectioned-pdf"
                and dst == fixture_root / "sectioned-pdf"
            )
            if is_second_publish:
                raise OSError("injected publish failure")
            if is_first_backup_restore:
                raise OSError("injected rollback failure")
            real_replace(source, destination)

        monkeypatch.setattr(generate_fixtures.os, "replace", fail_publish_then_rollback)
        captured_errors: list[generate_fixtures.RecoveryError] = []
        original_publish = generate_fixtures.publish_staged_artifacts

        def capture_recovery_error(stage_root: Path, publish_root: Path) -> None:
            try:
                original_publish(stage_root, publish_root)
            except generate_fixtures.RecoveryError as error:
                captured_errors.append(error)
                raise

        monkeypatch.setattr(
            generate_fixtures, "publish_staged_artifacts", capture_recovery_error
        )

        assert generate_fixtures.main() == 1
        assert len(captured_errors) == 1
        recovery_error = captured_errors[0]
        assert recovery_error.backup_root is not None
        assert recovery_error.stage_root is not None
        assert (recovery_error.backup_root / "sectioned-pdf" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "old-sectioned-pdf"
        assert (recovery_error.stage_root / "sectioned-pdf" / "payload.txt").read_text(
            encoding="utf-8"
        ) == "new-sectioned-pdf"
        captured_stderr = capsys.readouterr().err
        assert str(recovery_error.backup_root) in captured_stderr
        assert str(recovery_error.stage_root) in captured_stderr


class TestSourceDeterminism:
    """Source builders must produce byte-identical output on every call.

    ReportLab PDFs and python-docx DOCX files both embed timestamps and random
    IDs by default.  The production source builders must suppress this
    nondeterminism so that two generator runs produce identical source binaries,
    which in turn produces identical frozen docling_output.json.
    """

    def test_sectioned_pdf_is_byte_deterministic(self) -> None:
        """create_sectioned_pdf must return identical bytes on every call."""
        hashes = [
            hashlib.sha256(generate_fixtures.create_sectioned_pdf()).hexdigest()
            for _ in range(3)
        ]
        assert len(set(hashes)) == 1, (
            f"create_sectioned_pdf must be byte-deterministic, "
            f"got {len(set(hashes))} distinct hashes"
        )

    def test_complex_layout_pdf_is_byte_deterministic(self) -> None:
        """create_complex_layout_pdf must return identical bytes on every call."""
        hashes = [
            hashlib.sha256(generate_fixtures.create_complex_layout_pdf()).hexdigest()
            for _ in range(3)
        ]
        assert len(set(hashes)) == 1, (
            f"create_complex_layout_pdf must be byte-deterministic, "
            f"got {len(set(hashes))} distinct hashes"
        )

    def test_unicode_docx_is_byte_deterministic(self) -> None:
        """create_unicode_docx must return identical bytes on every call."""
        hashes = [
            hashlib.sha256(generate_fixtures.create_unicode_docx()).hexdigest()
            for _ in range(3)
        ]
        assert len(set(hashes)) == 1, (
            f"create_unicode_docx must be byte-deterministic, "
            f"got {len(set(hashes))} distinct hashes"
        )

    def test_docx_zip_entries_have_canonical_timestamps(self) -> None:
        """All ZIP entries in the DOCX must share a fixed canonical timestamp."""
        data = generate_fixtures.create_unicode_docx()
        z = zipfile.ZipFile(BytesIO(data))
        canonical = (1980, 1, 1, 0, 0, 0)
        for info in z.infolist():
            assert info.date_time == canonical, (
                f"ZIP entry {info.filename} has timestamp {info.date_time}, "
                f"expected canonical {canonical}"
            )

    def test_docx_remains_valid_after_canonicalization(self) -> None:
        """The canonicalized DOCX must still be openable by python-docx."""
        from docx import Document

        data = generate_fixtures.create_unicode_docx()
        doc = Document(BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        assert len(paragraphs) > 0, "Canonicalized DOCX must have readable content"

    def test_docx_canonicalization_preserves_entry_count_and_names(self) -> None:
        """Canonicalization must preserve entry count and filename order."""
        from docx import Document
        from io import BytesIO
        import zipfile

        # Get original ZIP entries from python-docx before canonicalization
        document = Document()
        document.add_heading("Test", level=1)
        document.add_paragraph("Content")
        raw_buf = BytesIO()
        document.save(raw_buf)
        raw_zip = zipfile.ZipFile(raw_buf)
        raw_names = [i.filename for i in raw_zip.infolist()]
        raw_zip.close()

        # Verify the canonicalization function itself preserves count/names
        canon_of_raw = generate_fixtures._canonicalize_docx_zip(raw_buf.getvalue())
        canon_of_raw_zip = zipfile.ZipFile(BytesIO(canon_of_raw))
        canon_of_raw_names = [i.filename for i in canon_of_raw_zip.infolist()]
        canon_of_raw_zip.close()

        assert (
            canon_of_raw_names == raw_names
        ), "Canonicalization must preserve entry names and ordering"

    def test_docx_canonicalization_preserves_compression_type(self) -> None:
        """Canonicalization must preserve each entry's compression type."""
        from docx import Document
        from io import BytesIO
        import zipfile

        document = Document()
        document.add_heading("Test", level=1)
        document.add_paragraph("Content")
        raw_buf = BytesIO()
        document.save(raw_buf)

        canon_data = generate_fixtures._canonicalize_docx_zip(raw_buf.getvalue())
        raw_zip = zipfile.ZipFile(raw_buf)
        canon_zip = zipfile.ZipFile(BytesIO(canon_data))

        for raw_info, canon_info in zip(raw_zip.infolist(), canon_zip.infolist()):
            assert raw_info.compress_type == canon_info.compress_type, (
                f"Compression type changed for {raw_info.filename}: "
                f"{raw_info.compress_type} -> {canon_info.compress_type}"
            )
        raw_zip.close()
        canon_zip.close()

    def test_docx_canonicalization_preserves_content_bytes(self) -> None:
        """Canonicalization must preserve each entry's decompressed content."""
        from docx import Document
        from io import BytesIO
        import zipfile

        document = Document()
        document.add_heading("Test", level=1)
        document.add_paragraph("Content")
        raw_buf = BytesIO()
        document.save(raw_buf)

        canon_data = generate_fixtures._canonicalize_docx_zip(raw_buf.getvalue())
        raw_zip = zipfile.ZipFile(raw_buf)
        canon_zip = zipfile.ZipFile(BytesIO(canon_data))

        for raw_info in raw_zip.infolist():
            raw_content = raw_zip.read(raw_info.filename)
            canon_content = canon_zip.read(raw_info.filename)
            assert (
                raw_content == canon_content
            ), f"Content changed for {raw_info.filename}"
        raw_zip.close()
        canon_zip.close()


class TestDoclingVersionGuard:
    """The generator must enforce docling==2.109.0 at runtime."""

    def test_guard_raises_on_wrong_docling_version(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The version guard must raise when docling != 2.109.0."""
        import importlib.metadata

        monkeypatch.setattr(
            importlib.metadata,
            "version",
            lambda pkg: (
                "2.93.0" if pkg == "docling" else importlib.metadata.version(pkg)
            ),
        )
        with pytest.raises(generate_fixtures.ValidationError, match="docling.*2.109.0"):
            generate_fixtures.assert_docling_version()

    def test_guard_passes_on_correct_docling_version(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The version guard must pass when docling == 2.109.0."""
        import importlib.metadata

        monkeypatch.setattr(
            importlib.metadata,
            "version",
            lambda pkg: (
                "2.109.0" if pkg == "docling" else importlib.metadata.version(pkg)
            ),
        )
        # Must not raise
        generate_fixtures.assert_docling_version()


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
